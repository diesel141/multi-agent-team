"""
pptx 内部 padding 検証スクリプト（汎用版）

検証項目:
- shape 内部 padding 対称性 (top_pad / bottom_pad / diff < 0.10 inch)
- inter-element gap 均等性 (max - min < 0.10 inch)
- 8pt baseline grid 整合 (gap が 0.111 inch 倍数か)

Designer 拡張:
- 同位置/包含 shape は装飾とみなして merge (pill rect + pill text / bullet circle + watch text 等)

Usage:
    python pptx_internal_padding.py <pptx_path> [--slides 17-22] [--container-min-w 1.5] [--container-min-h 3.0]

Exit code:
    0: 全 container PASS
    1: いずれかが FAIL
    2: pptx 読み込みエラー / 引数エラー

Origin:
    TASK-0007 Designer 白井美雪が `W:\\ai_training_for_rjc\\scripts\\verify_internal_padding.py`
    として起案。本リポに汎用化して取り込み（TASK-0010 chore）。
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation


EMU = 914400
GRID_8PT = 8.0 / 72.0  # 0.1111 inch
TOL = 0.025  # inch (約 1.8pt)


def is_container_factory(min_w: float, min_h: float):
    def is_container(sh):
        try:
            w = sh.width / EMU
            h = sh.height / EMU
            return w >= min_w and h >= min_h
        except Exception:
            return False

    return is_container


def inside(outer, inner):
    if outer is inner:
        return False
    try:
        ol, ot = outer.left, outer.top
        oe, ob = ol + outer.width, ot + outer.height
        il, it = inner.left, inner.top
        ie, ib = il + inner.width, it + inner.height
        return il > ol and it > ot and ie < oe and ib < ob
    except Exception:
        return False


def filter_top_level(children):
    """child が他の child の y range に完全包含されるなら装飾扱いで除外。"""
    result = []
    for c in children:
        contained = False
        for d in children:
            if d is c:
                continue
            try:
                c_top, c_bot = c.top, c.top + c.height
                d_top, d_bot = d.top, d.top + d.height
                if d_top <= c_top and d_bot >= c_bot and d.height > c.height:
                    contained = True
                    break
            except Exception:
                pass
        if not contained:
            result.append(c)
    return result


def merge_overlap_top(children):
    """top と高さがほぼ同じ shape は 1 論理要素扱い。"""
    out = []
    for c in children:
        merged = False
        for o in out:
            try:
                if (
                    abs(c.top - o.top) < TOL * EMU
                    and abs(c.height - o.height) < TOL * EMU
                ):
                    merged = True
                    break
            except Exception:
                pass
        if not merged:
            out.append(c)
    return out


def is_grid_aligned(g: float) -> bool:
    if g < 0.001:
        return True
    rem = g % GRID_8PT
    return rem < TOL or (GRID_8PT - rem) < TOL


def parse_slide_range(spec: str, total: int) -> list[int]:
    """e.g. "17-22" → [16,17,18,19,20,21] (0-indexed)。"all" → 全スライド。"""
    if spec == "all":
        return list(range(total))
    if "-" in spec:
        a, b = spec.split("-", 1)
        start, end = int(a), int(b)
        return list(range(start - 1, end))  # 1-indexed input → 0-indexed output
    # comma separated
    return [int(x) - 1 for x in spec.split(",")]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx_path", help="検証対象 .pptx のパス")
    parser.add_argument(
        "--slides",
        default="all",
        help='検証対象スライド範囲。例: "17-22"、"17,19,21"、"all"（既定）',
    )
    parser.add_argument("--container-min-w", type=float, default=1.5)
    parser.add_argument("--container-min-h", type=float, default=3.0)
    parser.add_argument(
        "--padding-diff-threshold",
        type=float,
        default=0.10,
        help="top_pad と bottom_pad の差の許容上限 (inch)",
    )
    parser.add_argument(
        "--gap-range-threshold",
        type=float,
        default=0.10,
        help="inter-element gap の max-min の許容上限 (inch)",
    )
    parser.add_argument(
        "--detail-slides",
        default=None,
        help='詳細 child ダンプを出力するスライド範囲（例: "18-19"）。既定はなし',
    )
    args = parser.parse_args()

    pptx_path = Path(args.pptx_path)
    if not pptx_path.exists():
        print(f"ERROR: {pptx_path} not found", file=sys.stderr)
        return 2

    try:
        prs = Presentation(str(pptx_path))
    except Exception as e:
        print(f"ERROR: failed to open {pptx_path}: {e}", file=sys.stderr)
        return 2

    target = parse_slide_range(args.slides, len(prs.slides))
    is_container = is_container_factory(args.container_min_w, args.container_min_h)

    print(
        "| Slide | container | size | top_pad | bottom_pad | diff | "
        "min_gap | max_gap | gap_range | grid8 | verdict |"
    )
    print("|---|---|---|---|---|---|---|---|---|---|---|")

    total_containers = 0
    fail_count = 0

    for idx in target:
        if idx < 0 or idx >= len(prs.slides):
            continue
        s = prs.slides[idx]
        containers = [sh for sh in s.shapes if is_container(sh)]
        for outer in containers:
            raw_children = [sh for sh in s.shapes if inside(outer, sh)]
            if not raw_children:
                continue
            children = filter_top_level(raw_children)
            children = merge_overlap_top(children)
            children.sort(key=lambda c: c.top)
            if not children:
                continue
            total_containers += 1

            top_pad = (children[0].top - outer.top) / EMU
            bottom_pad = (
                (outer.top + outer.height) - (children[-1].top + children[-1].height)
            ) / EMU
            diff = abs(top_pad - bottom_pad)
            gaps = [
                (children[i + 1].top - (children[i].top + children[i].height)) / EMU
                for i in range(len(children) - 1)
            ]
            min_gap = min(gaps) if gaps else 0
            max_gap = max(gaps) if gaps else 0
            gap_range = max_gap - min_gap
            grid_check = all(is_grid_aligned(g) for g in gaps) if gaps else True

            passes = (
                diff < args.padding_diff_threshold
                and gap_range < args.gap_range_threshold
            )
            verdict = "PASS" if passes else "FAIL"
            if not passes:
                fail_count += 1

            ow = outer.width / EMU
            oh = outer.height / EMU
            ol = outer.left / EMU
            ot = outer.top / EMU
            print(
                f"| {idx + 1} | ({ol:.2f},{ot:.2f}) | "
                f"{ow:.2f}x{oh:.2f} | {top_pad:.3f} | {bottom_pad:.3f} | "
                f"{diff:.3f} | {min_gap:.3f} | {max_gap:.3f} | {gap_range:.3f} | "
                f"{'OK' if grid_check else 'NG'} | {verdict} |"
            )

    print()
    print(
        f"SUMMARY: {total_containers - fail_count}/{total_containers} PASS "
        f"({fail_count} FAIL)"
    )

    if args.detail_slides:
        detail_target = parse_slide_range(args.detail_slides, len(prs.slides))
        print()
        print("---")
        print(f"詳細 child ダンプ (slides {args.detail_slides}):")
        for idx in detail_target:
            if idx < 0 or idx >= len(prs.slides):
                continue
            s = prs.slides[idx]
            print(f"\n=== Slide {idx + 1} ===")
            for outer in [sh for sh in s.shapes if is_container(sh)]:
                raw_children = [sh for sh in s.shapes if inside(outer, sh)]
                children = filter_top_level(raw_children)
                children = merge_overlap_top(children)
                children.sort(key=lambda c: c.top)
                print(
                    f"  container at ({outer.left / EMU:.2f}, "
                    f"{outer.top / EMU:.2f}) size "
                    f"{outer.width / EMU:.2f}x{outer.height / EMU:.2f} "
                    f"({len(children)} logical children, "
                    f"{len(raw_children)} raw)"
                )
                for c in children:
                    rel_top = (c.top - outer.top) / EMU
                    rel_end = rel_top + c.height / EMU
                    txt = ""
                    try:
                        if c.has_text_frame:
                            txt = c.text_frame.text[:30].replace("\n", " ")
                    except Exception:
                        pass
                    print(
                        f"    rel_y=[{rel_top:.3f}, {rel_end:.3f}] "
                        f"h={c.height / EMU:.3f} '{txt}'"
                    )

    return 0 if fail_count == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
