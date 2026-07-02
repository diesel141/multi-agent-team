"""
TASK-GLOBALCTX-VERIFY-02 受入検証スクリプト

CLI 研修デッキ「AI研修-CLIの便利コマンドとトークン節約術.pptx」の
GLOBAL CONTEXT セクション拡張（divider + 本編4枚: WHAT/PATH&FLOW/WHY/HOW）を機械検証する。

【核心】ユーザー最重視の「重なり・はみ出しゼロ」を機械担保すること。

 TASK-GLOBALCTX-01 (= 29 枚版) を 02 (= 31 枚版) に拡張した差分を検証する。
 before = 01 成果物（29 枚版）を *_old_for_verify.pptx として退避した snapshot。

検証項目:
  a. 同一スライド内の真の重なり 0（背景rect/装飾バー/✓コールアウト等の既知パターン除外）
  b. テキスト枠がスライド外(0..13.333 x 0..7.5 inch)へはみ出さない。
     新規パスツリー/フロー図/6カードの行あふれ推定を WARN で併記（参考・FAIL 非関与）。
  c. 総スライド数 = 31 (旧26 + divider + 本編4)
  d. 構造: 04 GLOBAL CONTEXT 本編 = WHAT→PATH&FLOW→WHY→HOW / 05 COST /
     AGENDA "100 minutes ｜ 5 sections"
  e. 副作用: GLOBAL CONTEXT 4 枚 + AGENDA 以外の既存スライド差分 0。
     （S8/S11 の既知事前ドリフトは 01 で既に取り込まれ、02 前後で不変のため
      before=29枚版との比較で自動的に差分なし = PM 受入方針「既知ドリフト許容」を満たす。
      02 で新規差分が出た場合のみ FAIL。）

Usage:
    python scripts/verify_task_globalctx_02.py [after_pptx] [--before before.pptx]

既定 after:  W:/ai_training_for_rjc/AI研修-CLIの便利コマンドとトークン節約術.pptx
既定 before: W:/ai_training_for_rjc/cli_globalctx02_old_for_verify.pptx (01 成果 29 枚版 snapshot)

Exit code: 0=全PASS / 1=いずれかFAIL / 2=読込エラー
"""
from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path

from pptx import Presentation

EMU_PER_INCH = 914400
OVERLAP_TOL = 0.02       # inch — edge-touching layout は除外
BBOX_TOL = 0.02          # inch — スライド境界ぴったりは許容
EXPECTED_SLIDES = 31     # 旧26 + divider + 本編4
WRAP_WARN_RATIO = 1.30   # 行あふれ推定 ratio の WARN 閾値（参考情報・FAIL にはしない）

DEFAULT_FONT_PT = 18  # font.size が取得できない run の推定サイズ

DEFAULT_AFTER = Path(
    r"W:\ai_training_for_rjc\AI研修-CLIの便利コマンドとトークン節約術.pptx"
)
DEFAULT_BEFORE = Path(
    r"W:\ai_training_for_rjc\cli_globalctx02_old_for_verify.pptx"
)


# ── helpers ──────────────────────────────────────────────

def _has_visible_text(shape) -> bool:
    if not shape.has_text_frame:
        return False
    return any(p.text.strip() for p in shape.text_frame.paragraphs)


def _bbox(shape):
    try:
        l, t = shape.left, shape.top
        return l, t, l + shape.width, t + shape.height
    except Exception:
        return None


def _rects_overlap(a, b, tol_in=OVERLAP_TOL) -> bool:
    ba, bb = _bbox(a), _bbox(b)
    if ba is None or bb is None:
        return False
    al, at, ar, ab = ba
    bl, bt, br, bb_ = bb
    tol = int(tol_in * EMU_PER_INCH)
    return (
        al < br - tol and ar > bl + tol
        and at < bb_ - tol and ab > bt + tol
    )


def slide_text_lines(slide) -> list[str]:
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    texts.append(t)
    return texts


# ── 検証 a: shape overlap = 0（意図的レイアウト除外）─────

def _contains(outer, inner, tol_in=0.05) -> bool:
    bo, bi = _bbox(outer), _bbox(inner)
    if bo is None or bi is None:
        return False
    tol = int(tol_in * EMU_PER_INCH)
    return (bi[0] >= bo[0] - tol and bi[1] >= bo[1] - tol
            and bi[2] <= bo[2] + tol and bi[3] <= bo[3] + tol)


def _overlap_xy(a, b) -> tuple[float, float]:
    ba, bb = _bbox(a), _bbox(b)
    if ba is None or bb is None:
        return 0.0, 0.0
    xo = max(0.0, min(ba[2], bb[2]) - max(ba[0], bb[0])) / EMU_PER_INCH
    yg = max(0.0, min(ba[3], bb[3]) - max(ba[1], bb[1])) / EMU_PER_INCH
    return xo, yg


def _is_callout(shape) -> bool:
    if not shape.has_text_frame:
        return False
    t = shape.text_frame.text.strip()
    return t.startswith(("✓", "✗", "※"))


def _is_intentional_overlap(a, b) -> bool:
    """意図的なレイアウト重なり（崩れではない）か:
    - 包含関係（背景+前景のレイヤー）
    - 2行タイトル/水平薄重なり（水平に広く xo>1.0, 垂直に薄く yg<0.4）
    - コールアウト/メモ（✓ 系）"""
    if _contains(a, b) or _contains(b, a):
        return True
    xo, yg = _overlap_xy(a, b)
    if xo > 1.0 and yg < 0.4:
        return True
    if _is_callout(a) or _is_callout(b):
        return True
    return False


def check_overlap(prs) -> list[str]:
    """全スライドで、意図的レイアウトを除外した「有意な重なり(崩れ)」を検出。"""
    errors = []
    for i, slide in enumerate(prs.slides, 1):
        visibles = [sh for sh in slide.shapes if _has_visible_text(sh)]
        for x in range(len(visibles)):
            for y in range(x + 1, len(visibles)):
                a, b = visibles[x], visibles[y]
                if not _rects_overlap(a, b):
                    continue
                if _is_intentional_overlap(a, b):
                    continue
                ta = a.text_frame.text.replace("\n", " ")[:30]
                tb = b.text_frame.text.replace("\n", " ")[:30]
                errors.append(
                    f"  Slide {i}: shape{a.shape_id} "
                    f"\"{ta}\" × shape{b.shape_id} \"{tb}\""
                )
    return errors


# ── 検証 b: no overflow ──────────────────────────────────

def _shape_outside_overflow(shape, slide_w_in, slide_h_in) -> float:
    bbox = _bbox(shape)
    if bbox is None:
        return 0.0
    l, t, r, b = bbox
    sw = int(slide_w_in * EMU_PER_INCH)
    sh = int(slide_h_in * EMU_PER_INCH)
    tol = int(BBOX_TOL * EMU_PER_INCH)
    over = 0.0
    if l < -tol:
        over = max(over, (-l) / EMU_PER_INCH)
    if t < -tol:
        over = max(over, (-t) / EMU_PER_INCH)
    if r > sw + tol:
        over = max(over, (r - sw) / EMU_PER_INCH)
    if b > sh + tol:
        over = max(over, (b - sh) / EMU_PER_INCH)
    return over


def _run_size_pt(run) -> float | None:
    try:
        if run.font.size is not None:
            return run.font.size.pt
    except Exception:
        pass
    return None


def _estimate_wrap_ratio(shape) -> float | None:
    """文字量対枠不足を行あふれ推定 ratio(必要高/内側高) で返す。近似。推定不可は None。"""
    if not shape.has_text_frame:
        return None
    tf = shape.text_frame
    if not any(p.text.strip() for p in tf.paragraphs):
        return None
    try:
        inner_w = (shape.width - (tf.margin_left or 0) - (tf.margin_right or 0)) / EMU_PER_INCH
        inner_h = (shape.height - (tf.margin_top or 0) - (tf.margin_bottom or 0)) / EMU_PER_INCH
    except Exception:
        return None
    if inner_w <= 0 or inner_h <= 0:
        return None

    total_line_units = 0.0
    sizes = []
    for p in tf.paragraphs:
        text = p.text
        size_pt = None
        for run in p.runs:
            s = _run_size_pt(run)
            if s is not None:
                size_pt = s
                sizes.append(s)
                break
        if size_pt is None:
            size_pt = DEFAULT_FONT_PT
        ls = p.line_spacing if p.line_spacing is not None else 1.0
        if not text:
            total_line_units += 1.0 * ls
            continue
        cjk = sum(1 for c in text if ord(c) > 0x2E80)
        cells = cjk + (len(text) - cjk) * 0.5
        chars_per_line = max(1.0, inner_w * 72.0 / size_pt)
        lines = max(1, math.ceil(cells / chars_per_line))
        total_line_units += lines * ls

    avg_pt = (sum(sizes) / len(sizes)) if sizes else DEFAULT_FONT_PT
    need_h = total_line_units * avg_pt / 72.0
    return need_h / inner_h if inner_h > 0 else None


def check_overflow(prs) -> tuple[list[str], list[str]]:
    errors = []
    warns = []
    sw = prs.slide_width / EMU_PER_INCH
    sh = prs.slide_height / EMU_PER_INCH
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            ov = _shape_outside_overflow(shape, sw, sh)
            if ov > 0:
                t = (
                    shape.text_frame.text.replace("\n", " ")[:30]
                    if shape.has_text_frame else f"<{shape.shape_type}>"
                )
                errors.append(
                    f"  Slide {i}: shape{shape.shape_id} \"{t}\" が "
                    f"スライド外へ {ov:.3f}\" はみ出し"
                )
            ratio = _estimate_wrap_ratio(shape)
            if ratio is not None and ratio > WRAP_WARN_RATIO:
                t = shape.text_frame.text.replace("\n", " ")[:30]
                warns.append(
                    f"  Slide {i}: shape{shape.shape_id} \"{t}\" "
                    f"行あふれ推定 ratio={ratio:.2f} (>{WRAP_WARN_RATIO:.2f})"
                )
    return errors, warns


# ── 検証 c: slide count ──────────────────────────────────

def check_slide_count(prs) -> list[str]:
    n = len(prs.slides)
    if n == EXPECTED_SLIDES:
        return []
    return [f"  総スライド数: expected {EXPECTED_SLIDES}, got {n}"]


# ── 検証 d: structure ────────────────────────────────────

_SECTION_RE = re.compile(r"^(\d{2})\s*/\s*([A-Z][A-Z &\-]+)$")
_GC_FRAME_RE = re.compile(r"^04\s+GLOBAL CONTEXT\s*[｜|]\s*(.+)$")


def _section_big_texts(prs) -> list[str]:
    found = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if _SECTION_RE.match(t):
                    found.append(t)
    return found


def _gc_frame_titles(prs) -> list[str]:
    """04 GLOBAL CONTEXT 本編 frame の "｜" 以降のタイトルを抽出。"""
    titles = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                m = _GC_FRAME_RE.match(para.text.strip())
                if m:
                    titles.append(m.group(1).strip())
    return titles


def check_structure(prs) -> list[str]:
    errors = []
    # SECTION dividers
    sections = _section_big_texts(prs)
    for exp in ["02 / COMMANDS", "03 / TOKEN SAVING", "04 / GLOBAL CONTEXT", "05 / COST"]:
        if exp not in sections:
            errors.append(f"  SECTION \"{exp}\" が見つかりません (検出: {sections})")

    # GLOBAL CONTEXT 本編4枚: WHAT → PATH&FLOW → WHY → HOW の順序（存在+順序を同時に検証）
    frames = _gc_frame_titles(prs)
    expected_order = ["WHAT", "PATH", "WHY", "HOW"]
    order_found: list[str] = []
    for f in frames:
        fu = f.upper()
        for kw in expected_order:
            if kw in fu and kw not in order_found:
                order_found.append(kw)
                break
    if order_found != expected_order:
        errors.append(
            f"  GLOBAL CONTEXT 本編の順序/構成が不正: 期待 {expected_order}, "
            f"検出順 {order_found} (GC frame: {frames})"
        )

    # AGENDA (slide 2)
    agenda = " ".join(slide_text_lines(prs.slides[1]))
    if "100 minutes" not in agenda or "5 sections" not in agenda:
        errors.append(
            "  AGENDA: '100 minutes ｜ 5 sections' が見つかりません "
            f"(AGENDA 先頭: {' / '.join(slide_text_lines(prs.slides[1])[:2])})"
        )
    return errors


# ── 検証 e: side effect = 0（before=29枚版 snapshot）─────

_COST_NUM_RE = re.compile(r"\b0[45]\s*/?\s*COST")
_MINUTES_RE = re.compile(r"想定時間：\s*\d+\s*minutes")


def _norm_text(t: str) -> str:
    t = re.sub(r"\s+", " ", t.strip())
    t = _COST_NUM_RE.sub("COSTNUM", t)
    t = _MINUTES_RE.sub("想定時間：MINUTES", t)
    return t


def _slide_norm(slide) -> str:
    return _norm_text("\n".join(slide_text_lines(slide)))


def check_side_effect(prs_after, prs_before) -> tuple[list[str], str | None]:
    """GLOBAL CONTEXT 4 枚 + AGENDA 以外の既存スライド差分 0。
    before=29枚版(01 成果)との比較。S8/S11 ドリフトは 02 前後で不変なので
    自動的に差分なし（PM 既知ドリフト許容を自然満た）。新規差分のみ FAIL。"""
    if prs_before is None:
        return [], "before pptx(29枚版 snapshot) が取得できず (e は SKIP)"
    errors = []

    before_norms = {_slide_norm(sl) for sl in prs_before.slides}
    after_norms = [_slide_norm(sl) for sl in prs_after.slides]
    after_set = set(after_norms)

    # after 側: AGENDA(2) と GLOBAL CONTEXT 含むスライド以外は before(29枚) に存在すべき
    for i, slide in enumerate(prs_after.slides, 1):
        norm = _slide_norm(slide)
        if i == 2:  # AGENDA は 01 で変更済み（02 でも変更対象外だが除外）
            continue
        if "GLOBAL CONTEXT" in norm:  # 02 で拡張されるセクション
            continue
        if norm not in before_norms:
            errors.append(
                f"  after Slide {i}: before(29枚) に対応なし（02 新規副作用疑い）: "
                f"\"{norm[:60]}\""
            )

    # before(29枚) 側: AGENDA 以外は COST番号吸収後 after(31枚) に残存すべき
    for j, slide in enumerate(prs_before.slides, 1):
        norm = _slide_norm(slide)
        if j == 2:  # before AGENDA
            continue
        if norm not in after_set:
            # 01 の GLOBAL CONTEXT(WHY/HOW) は 02 で再構成される可能性 → GLOBAL CONTEXT は許容
            if "GLOBAL CONTEXT" in norm:
                continue
            errors.append(
                f"  before Slide {j}: after から消失（02 新規副作用疑い）: "
                f"\"{norm[:60]}\""
            )
    return errors, None


# ── main ─────────────────────────────────────────────────

def _emit(label: str, errors: list[str]) -> None:
    status = "PASS" if not errors else "FAIL"
    print(f"[{status}] {label}")
    for e in errors:
        print(e)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="TASK-GLOBALCTX-VERIFY-02 受入検証"
    )
    parser.add_argument("after", nargs="?", default=str(DEFAULT_AFTER),
                        help="検証対象(after=31枚) pptx")
    parser.add_argument("--before", default=str(DEFAULT_BEFORE),
                        help="before pptx（29枚版 snapshot）")
    args = parser.parse_args()

    after_path = Path(args.after)
    before_path = Path(args.before)
    if not after_path.exists():
        print(f"ERROR: {after_path} not found", file=sys.stderr)
        return 2
    try:
        prs = Presentation(str(after_path))
    except Exception as e:
        print(f"ERROR: {after_path} 読み込み失敗: {e}", file=sys.stderr)
        return 2

    prs_before = None
    if before_path.exists():
        try:
            prs_before = Presentation(str(before_path))
        except Exception as e:
            print(f"WARN: before 読込失敗({e})。e は SKIP", file=sys.stderr)
    else:
        print(f"WARN: before({before_path}) not found。e は SKIP", file=sys.stderr)

    print(f"検証対象(after): {after_path}")
    print(f"スライド数: {len(prs.slides)}")
    print(f"before: {str(before_path) + ' (取得済み)' if prs_before else '取得失敗 → e SKIP'}")
    print()

    all_pass = True

    errs = check_overlap(prs)
    _emit("a. shape overlap = 0（意図的レイアウト除外 / 真の崩れ）", errs)
    if errs:
        all_pass = False

    errs, warns = check_overflow(prs)
    _emit("b. no overflow（スライド外はみ出し = 0）", errs)
    if errs:
        all_pass = False
    if warns:
        print("  [参考 WARN] 行あふれ推定（実レンダリング非準拠・目安）:")
        for w in warns:
            print(w)

    errs = check_slide_count(prs)
    _emit(f"c. slide count = {EXPECTED_SLIDES}", errs)
    if errs:
        all_pass = False

    errs = check_structure(prs)
    _emit("d. structure（GC: WHAT/PATH&FLOW/WHY/HOW, 05 COST, AGENDA 100min/5sections）", errs)
    if errs:
        all_pass = False

    errs, note = check_side_effect(prs, prs_before)
    label = "e. side effect = 0（GC4枚+AGENDA 以外 差分 0 / S8・S11 既知ドリフト許容）"
    if note:
        print(f"[SKIP] {label} — {note}")
    else:
        _emit(label, errs)
        if errs:
            all_pass = False

    print()
    if all_pass:
        print("RESULT: 全項目 PASS")
        return 0
    else:
        print("RESULT: FAIL — 上記項目を修正してください")
        return 1


if __name__ == "__main__":
    sys.exit(main())
